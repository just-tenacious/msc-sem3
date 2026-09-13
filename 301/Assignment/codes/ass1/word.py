from docx import Document

doc = Document()

doc.add_heading("Assignment 1: Reading, Cleaning and Pre-processing CSV Data", level=1)

doc.add_paragraph(
    "Topics Covered: Reading & Writing CSV, Data Cleaning using Pandas and NumPy"
)

questions = [
    (
        "1. Read a CSV file and display the first 5 rows",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

print(df.head())"""
    ),
    (
        "2. Identify missing values in each column",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

missing = df.isnull().sum()

print("Missing values in each column:")
print(missing)"""
    ),
    (
        "3. Fill missing numeric values with column mean",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

df["Age"] = df["Age"].fillna(df["Age"].mean())

print(df)"""
    ),
    (
        "4. Fill missing categorical values with mode",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

df["Gender"] = df["Gender"].fillna(df["Gender"].mode()[0])

print(df)"""
    ),
    (
        "5. Remove duplicate records from the dataset",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

df = df.drop_duplicates()

print(df)"""
    ),
    (
        "6. Convert a column datatype (Age to int)",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

df["Age"] = df["Age"].fillna(df["Age"].mean()).astype(int)

print(df.dtypes)
print(df)"""
    ),
    (
        "7. Replace all negative salary values with median salary",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

median_salary = df["Salary"][df["Salary"] > 0].median()

df["Salary"] = df["Salary"].apply(
    lambda x: median_salary if x < 0 else x
)

print(df)"""
    ),
    (
        "8. Rename columns for better readability",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

df.rename(
    columns={
        "EmpName": "Employee_Name",
        "Dept": "Department"
    },
    inplace=True
)

print(df.columns)
print(df.head())"""
    ),
    (
        "9. Reorder columns in a desired sequence",
        """import pandas as pd

df = pd.read_csv("./data/data1.csv")

df.rename(
    columns={
        "EmpName": "Employee_Name",
        "Dept": "Department"
    },
    inplace=True
)

df = df[
    [
        "Employee_ID",
        "Employee_Name",
        "Age",
        "Gender",
        "Salary",
        "Department"
    ]
]

print(df.head())"""
    ),
    (
        "10. Use NumPy to replace all NaN values in a 2D array with zero",
        """import numpy as np

arr = np.array([
    [1, np.nan, 3],
    [4, 5, np.nan]
])

arr = np.nan_to_num(arr)

print(arr)"""
    )
]


for question, code in questions:
    doc.add_heading(question, level=2)
    doc.add_paragraph("Python Code:")
    doc.add_paragraph(code)


doc.save("Assignment_1_Code.docx")

print("Word file created successfully.")